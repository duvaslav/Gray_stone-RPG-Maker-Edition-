#!/usr/bin/env python3
"""Extract the Gray Stone design workbook into spec/*.json.

The XLSX is the human-facing source specification. It is NOT a runtime database:
nothing in the shipped game reads it. This script is the only bridge — it turns
each sheet into a list-of-rows JSON file under spec/, which the build-time
generators in tools/build/ consume.

Usage:  python3 tools/extract_workbook.py <path-to.xlsx>
"""
import json, os, re, sys

def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    try:
        import openpyxl
    except ImportError:
        print("openpyxl required:  pip install openpyxl", file=sys.stderr)
        return 2

    src = sys.argv[1]
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out = os.path.join(root, "spec")
    os.makedirs(out, exist_ok=True)

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    index = {}
    for name in wb.sheetnames:
        rows = []
        for raw in wb[name].iter_rows(values_only=True):
            if raw is None:
                continue
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in raw):
                continue
            rows.append(["" if c is None else str(c) for c in raw])
        # drop trailing all-empty columns so diffs stay small
        if rows:
            width = max(len(r) for r in rows)
            while width > 0 and all(len(r) < width or r[width - 1] == "" for r in rows):
                width -= 1
            rows = [r[:width] for r in rows]
        slug = re.sub(r"[^A-Za-z0-9_]", "_", name)
        with open(os.path.join(out, slug + ".json"), "w", encoding="utf-8") as fh:
            json.dump(rows, fh, ensure_ascii=False, indent=0)
        index[name] = {"file": slug + ".json", "rows": len(rows)}

    with open(os.path.join(out, "_index.json"), "w", encoding="utf-8") as fh:
        json.dump(index, fh, ensure_ascii=False, indent=1)
    print(f"extracted {len(index)} sheets -> spec/")
    return 0

if __name__ == "__main__":
    sys.exit(main())
